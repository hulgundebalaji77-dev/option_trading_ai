import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

def generate_excel_report(trades_df: pd.DataFrame, initial_capital=100000.0, output_name=None):
    if trades_df.empty:
        print("⚠️ कोणताही ट्रेड डेटा उपलब्ध नाही.")
        return None

    if not output_name:
        today = datetime.datetime.now().strftime("%Y-%m-%d")
        output_name = f"AI_Option_Report_{today}.xlsx"

    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_log = wb.create_sheet(title="Trade Log")
    ws_log.views.sheetView[0].showGridLines = True

    # Styles
    font_white = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_win = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_loss = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    border = Border(left=Side(style='thin', color='D9D9D9'),
                    right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'),
                    bottom=Side(style='thin', color='D9D9D9'))

    # Metrics
    total_trades = len(trades_df)
    wins = trades_df[trades_df["Result"] == "TARGET_HIT"]
    losses = trades_df[trades_df["Result"] == "SL_HIT"]
    win_rate = (len(wins) / total_trades) * 100 if total_trades > 0 else 0
    total_pnl = trades_df["Net P&L (₹)"].sum()

    ws_sum["B2"] = "AI TRADING DAILY SUMMARY"
    ws_sum["B2"].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")

    metrics = [
        ("Initial Capital", initial_capital),
        ("Final Balance", initial_capital + total_pnl),
        ("Net P&L (₹)", total_pnl),
        ("Total Trades", total_trades),
        ("Win Rate (%)", f"{win_rate:.2f}%"),
        ("Winning Trades", len(wins)),
        ("Losing Trades", len(losses))
    ]

    for i, (k, v) in enumerate(metrics, start=4):
        ws_sum[f"B{i}"] = k
        ws_sum[f"C{i}"] = v
        ws_sum[f"B{i}"].font = Font(name="Calibri", size=10, bold=True)
        ws_sum[f"C{i}"].font = Font(name="Calibri", size=10)
        ws_sum[f"B{i}"].border = border
        ws_sum[f"C{i}"].border = border

    # Log Sheet
    headers = list(trades_df.columns)
    ws_log.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws_log.cell(row=1, column=c)
        cell.font = font_white
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(trades_df.itertuples(index=False), start=2):
        for c_idx, val in enumerate(row, start=1):
            c = ws_log.cell(row=r_idx, column=c_idx, value=val)
            c.border = border
            if headers[c_idx - 1] == "Result":
                c.fill = fill_win if val == "TARGET_HIT" else fill_loss

    for ws in [ws_sum, ws_log]:
        for col in ws.columns:
            w = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(w + 3, 12)

    wb.save(output_name)
    print(f"📊 एक्सेल रिपोर्ट तयार झाली: {output_name}")
    return output_name
