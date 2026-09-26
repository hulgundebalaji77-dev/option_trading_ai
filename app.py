# app.py (सुरुवातीच्या ओळी)

import re
from datetime import datetime, timedelta
import feedparser
import nltk
from nltk.sentiment.vader import SentimentIntensityAnalyzer
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import pyotp
import requests
from SmartApi import SmartConnect
import streamlit as st  # <--- हे st वापरण्याआधी असणे आवश्यक आहे

from config import CONFIG

# ==========================================
# Angel One Data Fetcher (इम्पोर्ट्सच्या खाली)
# ==========================================
@st.cache_data(ttl=300)
def load_angel_data():
    try:
        api_key = str(CONFIG.get("ANGEL_API_KEY", "")).strip()
        client_code = str(CONFIG.get("ANGEL_CLIENT_CODE", "")).strip()
        password = str(CONFIG.get("ANGEL_PASSWORD", "")).strip()
        totp_raw = str(CONFIG.get("ANGEL_TOTP_TOKEN", "")).strip()

        totp_token = re.sub(r'[^A-Za-z2-7]', '', totp_raw).upper()

        if not all([api_key, client_code, password, totp_token]):
            st.sidebar.error("❌ क्रेडेंशियल्स अपूर्ण आहेत किंवा चुकीचे फॉरमॅट आहे.")
            return None

        totp = pyotp.TOTP(totp_token).now()

        smart_api = SmartConnect(api_key=api_key)
        session_data = smart_api.generateSession(client_code, password, totp)

        if not session_data.get("status"):
            st.sidebar.error(f"❌ लॉगिन अयशस्वी: {session_data.get('message')}")
            return None

        to_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        from_date = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d 09:15")

        historic_param = {
            "exchange": "NSE",
            "symboltoken": "99926000",
            "interval": "FIVE_MINUTE",
            "fromdate": from_date,
            "todate": to_date
        }
        res = smart_api.getCandleData(historic_param)
        
        if res and res.get("status") and res.get("data"):
            df = pd.DataFrame(res["data"], columns=["timestamp", "open", "high", "low", "close", "volume"])
            df["timestamp"] = pd.to_datetime(df["timestamp"])
            return df
        return None

    except Exception as e:
        st.sidebar.error(f"Angel One एरर: {e}")
        return None
        import re

@st.cache_data(ttl=300)
def load_angel_data():
    try:
        # सर्व क्रेडेंशियल्समधून अतिरिक्त स्पेसेस काढून घेणे
        api_key = str(CONFIG.get("ANGEL_API_KEY", "")).strip()
        client_code = str(CONFIG.get("ANGEL_CLIENT_CODE", "")).strip()
        password = str(CONFIG.get("ANGEL_PASSWORD", "")).strip()
        totp_raw = str(CONFIG.get("ANGEL_TOTP_TOKEN", "")).strip()

        # TOTP की मधील अनावश्यक स्पेसेस काढून फक्त Base32 कॅरेक्टर्स ठेवणे
        totp_token = re.sub(r'[^A-Za-z2-7]', '', totp_raw).upper()

        if not all([api_key, client_code, password, totp_token]):
            st.sidebar.error("❌ क्रेडेंशियल्स अपूर्ण आहेत किंवा चुकीचे फॉरमॅट आहे.")
            return None

        # TOTP जनरेट करणे
        totp = pyotp.TOTP(totp_token).now()

        smart_api = SmartConnect(api_key=api_key)
        session_data = smart_api.generateSession(client_code, password, totp)

        if not session_data.get("status"):
            st.sidebar.error(f"❌ लॉगिन अयशस्वी: {session_data.get('message')}")
            return None

        # हिस्टॉरिकल डेटा फेच करणे
        to_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        from_date = (datetime.now() - timedelta(days=7)).strftime("%Y-%m-%d 09:15")

        historic_param = {
            "exchange": "NSE",
            "symboltoken": "99926000",
            "interval": "FIVE_MINUTE",
            "fromdate": from_date,
            "todate": to_date
        }
        res = smart_api.getCandleData(historic_param)
        
        if res and res.get("status") and res.get("data"):
            df = pd.DataFrame(res["data"], columns=["timestamp", "open", "high", "low", "close", "volume"])
            df["timestamp"] = pd.to_datetime(df["timestamp"])
            return df
        return None

    except Exception as e:
        st.sidebar.error(f"Angel One एरर: {e}")
        return None
    from datetime import datetime, timedelta
import feedparser
import nltk
from nltk.sentiment.vader import SentimentIntensityAnalyzer
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import pyotp
import requests
from SmartApi import SmartConnect
import streamlit as st

# कॉन्फिग फाइल इम्पोर्ट
from config import CONFIG

# ==========================================
# 1. Angel One Data Fetcher
# ==========================================
@st.cache_data(ttl=300)
def load_angel_data():
    try:
        smart_api = SmartConnect(api_key=CONFIG["ANGEL_API_KEY"])
        totp = pyotp.TOTP(CONFIG["ANGEL_TOTP_TOKEN"]).now()
        smart_api.generateSession(CONFIG["ANGEL_CLIENT_CODE"], CONFIG["ANGEL_PASSWORD"], totp)

        to_date = datetime.now().strftime("%Y-%m-%d %H:%M")
        from_date = (datetime.now() - timedelta(days=5)).strftime("%Y-%m-%d 09:15")

        historic_param = {
            "exchange": "NSE",
            "symboltoken": "99926000",  # Nifty 50 Index
            "interval": "FIVE_MINUTE",
            "fromdate": from_date,
            "todate": to_date
        }
        res = smart_api.getCandleData(historic_param)
        if res and res.get("status"):
            df = pd.DataFrame(res["data"], columns=["timestamp", "open", "high", "low", "close", "volume"])
            df["timestamp"] = pd.to_datetime(df["timestamp"])
            return df
        return None
    except Exception as e:
        st.error(f"Angel One एरर: {e}")
        return None

# NLTK Lexicon Setup
try:
    nltk.data.find('sentiment/vader_lexicon.zip')
except LookupError:
    nltk.download('vader_lexicon', quiet=True)

# ==========================================
# 2. Risk Engine
# ==========================================
class DynamicRiskEngine:
    def __init__(self, rr_ratio=2.0):
        self.rr_ratio = rr_ratio

    def get_trade_levels(self, price, atr, action):
        atr_val = 25.0 if (atr is None or atr <= 0 or np.isnan(atr)) else float(atr)
        sl_points = round(max(atr_val * 1.5, 20.0), 1)
        tgt_points = round(sl_points * self.rr_ratio, 1)

        if action == "BUY_CE":
            sl = round(price - sl_points, 1)
            tgt = round(price + tgt_points, 1)
        elif action == "BUY_PE":
            sl = round(price + sl_points, 1)
            tgt = round(price - tgt_points, 1)
        else:
            return {}

        return {
            "entry": round(price, 1),
            "sl": sl,
            "target": tgt,
            "sl_pts": sl_points,
            "tgt_pts": tgt_points
        }

# ==========================================
# 3. Feature Extraction
# ==========================================
def extract_features(df):
    f_df = pd.DataFrame(index=df.index)
    f_df['returns'] = df['close'].pct_change()
    f_df['hl_spread'] = (df['high'] - df['low']) / df['close']

    ema9 = df['close'].ewm(span=9).mean()
    ema21 = df['close'].ewm(span=21).mean()
    f_df['ema_ratio'] = (ema9 / ema21) - 1.0

    delta = df['close'].diff()
    gain = (delta.where(delta > 0, 0)).rolling(window=14).mean()
    loss = (-delta.where(delta < 0, 0)).rolling(window=14).mean()
    rs = gain / (loss + 1e-9)
    f_df['rsi'] = 100 - (100 / (1 + rs))

    tr1 = df['high'] - df['low']
    tr2 = (df['high'] - df['close'].shift()).abs()
    tr3 = (df['low'] - df['close'].shift()).abs()
    f_df['atr'] = pd.concat([tr1, tr2, tr3], axis=1).max(axis=1).rolling(14).mean()
    return f_df.replace([np.inf, -np.inf], np.nan).fillna(0.0)

# ==========================================
# 4. ML AI Engine (Softmax Classifier)
# ==========================================
class OptionAIEngineV2:
    def __init__(self):
        self.weights = None
        self.bias = None
        self.mean = None
        self.std = None

    def train(self, df, features):
        future_diff = df['close'].shift(-3) - df['close']
        conditions = [future_diff >= 20, future_diff <= -20]
        targets = np.select(conditions, [1, 2], default=0)

        clean_df = df.copy()
        clean_df['target'] = targets
        clean_df = clean_df.dropna()

        X = features.loc[clean_df.index].values.astype(np.float64)
        y = clean_df['target'].values.astype(np.int64)

        n_features = X.shape[1]
        self.weights = np.zeros((3, n_features))
        self.bias = np.zeros(3)

        self.mean = np.mean(X, axis=0)
        self.std = np.std(X, axis=0) + 1e-7
        X_norm = (X - self.mean) / self.std

        lr = 0.05
        for _ in range(60):
            scores = np.dot(X_norm, self.weights.T) + self.bias
            exp_scores = np.exp(scores - np.max(scores, axis=1, keepdims=True))
            probs = exp_scores / np.sum(exp_scores, axis=1, keepdims=True)

            y_onehot = np.zeros_like(probs)
            np.put_along_axis(y_onehot, y[:, None], 1, axis=1)

            grad = probs - y_onehot
            dW = np.dot(grad.T, X_norm) / len(X)
            db = np.sum(grad, axis=0) / len(X)

            self.weights -= lr * dW
            self.bias -= lr * db

    def predict(self, feature_row):
        try:
            x_raw = np.asarray(feature_row.values, dtype=np.float64).reshape(1, -1)
            x_norm = (x_raw - self.mean) / self.std
            scores = np.dot(x_norm, self.weights.T) + self.bias
            exp_scores = np.exp(scores - np.max(scores, axis=1, keepdims=True))
            probs = (exp_scores / np.sum(exp_scores, axis=1, keepdims=True))[0]

            pred_class = int(np.argmax(probs))
            confidence = float(probs[pred_class])
            action_map = {0: "HOLD", 1: "BUY_CE", 2: "BUY_PE"}
            return action_map.get(pred_class, "HOLD"), confidence
        except Exception:
            return "HOLD", 0.0

# ==========================================
# 5. Excel Report Generator
# ==========================================
def create_excel_report(trades_df, capital, filename="AI_Option_Report.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trade Summary"
    fill_hdr = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    headers = list(trades_df.columns)
    ws.append(headers)
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.fill = fill_hdr
        cell.font = font_hdr
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(trades_df.itertuples(index=False), start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(horizontal="center")

    wb.save(filename)
    return filename

# ==========================================
# 6. Streamlit UI Layout
# ==========================================
st.set_page_config(page_title="AI Option Trading", page_icon="📈", layout="wide")
st.title("🤖 Nifty 50 AI Option Trading & Backtest System")

# API Status in Sidebar
st.sidebar.header("🔌 API Status")
angel_df = load_angel_data()
if angel_df is not None and not angel_df.empty:
    st.sidebar.success(f"Angel One Connected ({len(angel_df)} candles)")
else:
    st.sidebar.warning("Angel One Not Connected (Using Simulation Data)")

st.sidebar.header("⚙️ Strategy Parameters")
capital = st.sidebar.number_input("Starting Capital (₹)", value=100000, step=10000)
lot_size = st.sidebar.number_input("Lot Size (Nifty)", value=50, step=25)
rr_ratio = st.sidebar.slider("Risk-Reward Ratio", 1.0, 3.0, 2.0, 0.5)
min_conf = st.sidebar.slider("Minimum AI Confidence", 0.40, 0.90, 0.55, 0.05)

tab1, tab2, tab3 = st.tabs(["📊 Backtest Engine", "📰 Live News Sentiment", "📁 Export Reports"])

with tab1:
    st.subheader("Historical 5-Min Strategy Backtest")
    
    use_live_data = st.checkbox("Angel One चा खरा डेटा वापरा (उपलब्ध असल्यास)", value=False)
    
    if st.button("🚀 Run Backtest Now", type="primary"):
        with st.spinner("डेटा प्रोसेस आणि AI मॉडेल रन होत आहे..."):
            if use_live_data and angel_df is not None and len(angel_df) >= 200:
                df = angel_df.copy().rename(columns={"timestamp": "datetime"})
            else:
                np.random.seed(42)
                prices = np.cumprod(1 + np.random.normal(0.0001, 0.0025, 1200)) * 24500
                df = pd.DataFrame({
                    "datetime": pd.date_range("2026-06-01 09:15", periods=1200, freq="5min"),
                    "open": prices,
                    "high": prices * 1.002,
                    "low": prices * 0.998,
                    "close": prices,
                    "volume": np.random.randint(1000, 15000, 1200)
                })

            features = extract_features(df)
            split_idx = int(len(df) * 0.70)
            train_df = df.iloc[:split_idx]
            train_feats = features.iloc[:split_idx]

            ai = OptionAIEngineV2()
            ai.train(train_df, train_feats)

            risk_mgr = DynamicRiskEngine(rr_ratio=rr_ratio)
            test_df = df.iloc[split_idx:]
            trades = []
            in_trade = False
            trade_type, entry_p, sl_p, tgt_p = None, 0, 0, 0
            curr_cap = float(capital)
            t_id = 1

            for i in range(len(test_df) - 1):
                curr = test_df.iloc[i]
                feat = features.loc[[test_df.index[i]]]

                if in_trade:
                    h, l = curr['high'], curr['low']
                    exit_p, res = None, None
                    if trade_type == "BUY_CE":
                        if l <= sl_p:
                            exit_p, res = sl_p, "SL_HIT"
                        elif h >= tgt_p:
                            exit_p, res = tgt_p, "TARGET_HIT"
                    elif trade_type == "BUY_PE":
                        if h >= sl_p:
                            exit_p, res = sl_p, "SL_HIT"
                        elif l <= tgt_p:
                            exit_p, res = tgt_p, "TARGET_HIT"

                    if exit_p:
                        pts = (exit_p - entry_p) if trade_type == "BUY_CE" else (entry_p - exit_p)
                        pnl = round(pts * lot_size, 2)
                        curr_cap += pnl
                        trades.append({
                            "Trade ID": f"TRD_{t_id:03d}",
                            "Action": trade_type,
                            "Entry Price": entry_p,
                            "Exit Price": exit_p,
                            "Result": res,
                            "Points": round(pts, 2),
                            "Net P&L (₹)": pnl,
                            "Capital Balance (₹)": round(curr_cap, 2)
                        })
                        t_id += 1
                        in_trade = False

                if not in_trade:
                    action, conf = ai.predict(feat)
                    if action in ["BUY_CE", "BUY_PE"] and conf >= min_conf:
                        atr_v = feat['atr'].values[0]
                        levels = risk_mgr.get_trade_levels(curr['close'], atr_v, action)
                        trade_type = action
                        entry_p, sl_p, tgt_p = levels['entry'], levels['sl'], levels['target']
                        in_trade = True

            st.session_state['trades_df'] = pd.DataFrame(trades)
            st.success("✅ बॅकटेस्ट यशस्वीरित्या पूर्ण झाले!")

    # रिझल्ट्स डिस्प्ले
    if 'trades_df' in st.session_state:
        tdf = st.session_state['trades_df']
        if not tdf.empty:
            wins = tdf[tdf["Result"] == "TARGET_HIT"]
            win_rate = (len(wins) / len(tdf)) * 100
            total_pnl = tdf["Net P&L (₹)"].sum()

            st.markdown("---")
            st.subheader("📊 बॅकटेस्ट निकाल (Results)")
            col1, col2, col3, col4 = st.columns(4)
            col1.metric("Total Trades", len(tdf))
            col2.metric("Win Rate", f"{win_rate:.1f}%")
            col3.metric("Net P&L (₹)", f"₹{total_pnl:,.2f}", delta=f"{total_pnl:,.2f}")
            col4.metric("Ending Capital", f"₹{tdf['Capital Balance (₹)'].iloc[-1]:,.2f}")

            st.subheader("📈 Capital Growth (Equity Curve)")
            st.line_chart(tdf.set_index("Trade ID")["Capital Balance (₹)"])
            
            st.subheader("📋 Trade Logs")
            st.dataframe(tdf, use_container_width=True)
        else:
            st.warning("या कालावधीत निकष पूर्ण न झाल्यामुळे एकही ट्रेड मिळाला नाही.")

with tab2:
    st.subheader("Market News & Sentiment Radar")
    if st.button("🔄 Check Live News Sentiment"):
        with st.spinner("ताज्या बातम्या वाचल्या जात आहेत..."):
            try:
                headers = {"User-Agent": "Mozilla/5.0"}
                res = requests.get(
                    "https://news.google.com/rss/search?q=Nifty+50+Stock+Market+India&hl=en-IN&gl=IN&ceid=IN:en", 
                    headers=headers, 
                    timeout=10
                )
                feed = feedparser.parse(res.content)
                sia = SentimentIntensityAnalyzer()
                scores = [sia.polarity_scores(e.title)['compound'] for e in feed.entries[:8] if hasattr(e, 'title')]
                avg = sum(scores) / len(scores) if scores else 0.0

                sentiment = "BULLISH" if avg >= 0.10 else ("BEARISH" if avg <= -0.10 else "NEUTRAL")
                c1, c2 = st.columns(2)
                c1.metric("Overall Sentiment", sentiment)
                c2.metric("Sentiment Score", round(avg, 3))
                if sentiment == "BULLISH":
                    st.success("बाजारात तेजीचे संकेत आहेत (CE Bias).")
                elif sentiment == "BEARISH":
                    st.error("बाजारात मंदीचे संकेत आहेत (PE Bias).")
                else:
                    st.info("बाजार सध्या न्यूट्रल/साइडवेज आहे.")
            except Exception as e:
                st.warning(f"न्यूज लोड करताना अडचण: {e}")

with tab3:
    st.subheader("Download Formatted Excel Report")
    if 'trades_df' in st.session_state and not st.session_state['trades_df'].empty:
        fname = create_excel_report(st.session_state['trades_df'], capital)
        with open(fname, "rb") as f:
            st.download_button(
                "📥 Download Excel P&L Report", 
                f, 
                file_name=fname, 
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    else:
        st.warning("आधी बॅकटेस्ट रन करा, त्यानंतर एक्सेल डाउनलोड करता येईल.")
